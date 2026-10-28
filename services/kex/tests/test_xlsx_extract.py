"""Excel extraction: cached values first, formulas as fallback, legacy .xls routed to xlrd.

The Asgard failure of 2026-09-07 ("Excel file contained no data" on a financial model) was a
workbook whose formulas carried no cached results — data_only=True saw only None.
"""
import io
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sources.file_handler import _extract_xlsx, extract_text  # noqa: E402


def _workbook_bytes(fill) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    fill(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_values_are_extracted_normally():
    data = _workbook_bytes(lambda ws: ws.append(["Revenue", 1200, "EUR"]))
    text = _extract_xlsx(data)
    assert "Sheet: Model" in text
    assert "Revenue | 1200 | EUR" in text
    assert "formulas without cached results" not in text


def test_formula_only_workbook_falls_back_to_formulas():
    def fill(ws):
        ws["A1"] = "=SUM(B2:B13)"
        ws["A2"] = "=A1*1.19"

    data = _workbook_bytes(fill)  # openpyxl writes formulas WITHOUT cached values
    text = _extract_xlsx(data)
    assert "formulas without cached results" in text
    assert "=SUM(B2:B13)" in text
    assert "=A1*1.19" in text


def test_truly_empty_workbook_still_raises():
    data = _workbook_bytes(lambda ws: None)
    with pytest.raises(ValueError, match="contained no data"):
        _extract_xlsx(data)


def test_legacy_xls_is_routed_to_xlrd_not_rejected():
    # Not a real .xls body — the point is the ROUTE: .xls must reach xlrd (which then
    # complains about the bytes) instead of the "convert to .xlsx" rejection.
    with pytest.raises(Exception) as exc:
        extract_text(b"not-an-xls", "application/octet-stream", "legacy.xls")
    assert "convert to the modern format" not in str(exc.value)


# Asgard 2026-10-28: four "Datenkatalog (altes Format)" uploads failed with openpyxl's
# "Value must be either numerical or a string containing a wildcard" — a filter criterion
# the library rejects while loading. The workbook is read again without its autofilters.
def _filtered_workbook_bytes() -> bytes:
    def fill(ws):
        ws.append(["Spalte", "Wert"])
        ws.append(["Katalog", 42])
        ws.auto_filter.ref = "A1:B2"
        ws.auto_filter.add_filter_column(0, ["Katalog"])

    return _workbook_bytes(fill)


def test_autofilters_are_stripped_and_cells_survive():
    import zipfile

    from sources.file_handler import _xlsx_without_autofilters

    data = _filtered_workbook_bytes()
    assert any(b"<autoFilter" in zipfile.ZipFile(io.BytesIO(data)).read(n) for n in zipfile.ZipFile(io.BytesIO(data)).namelist())
    stripped = _xlsx_without_autofilters(data)
    for name in zipfile.ZipFile(io.BytesIO(stripped)).namelist():
        assert b"<autoFilter" not in zipfile.ZipFile(io.BytesIO(stripped)).read(name)
    assert "Katalog | 42" in _extract_xlsx(stripped)


def test_wildcard_rejection_falls_back_to_a_workbook_without_filters(monkeypatch):
    import openpyxl

    real = openpyxl.load_workbook
    calls = []

    def strict(fh, **kw):
        raw = fh.getvalue()
        calls.append(b"<autoFilter" in zipfile_has(raw))
        if calls[-1]:
            raise ValueError("Value must be either numerical or a string containing a wildcard")
        return real(io.BytesIO(raw), **kw)

    def zipfile_has(raw):
        import zipfile

        z = zipfile.ZipFile(io.BytesIO(raw))
        return b"".join(z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/"))

    monkeypatch.setattr(openpyxl, "load_workbook", strict)
    text = _extract_xlsx(_filtered_workbook_bytes())
    assert "Katalog | 42" in text
    assert calls == [True, False], "first load rejected, retry without filters accepted"


def test_other_value_errors_still_surface(monkeypatch):
    import openpyxl

    def broken(fh, **kw):
        raise ValueError("something else entirely")

    monkeypatch.setattr(openpyxl, "load_workbook", broken)
    with pytest.raises(ValueError, match="something else"):
        _extract_xlsx(_filtered_workbook_bytes())
